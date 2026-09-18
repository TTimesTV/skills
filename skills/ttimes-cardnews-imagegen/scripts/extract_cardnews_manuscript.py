#!/usr/bin/env python3
"""Extract TTimes card-news manuscript structure from a DOCX.

Usage:
    python extract_cardnews_manuscript.py input.docx -o manuscript.json

Observed grammar:
- leading PXX. paragraphs before <1> are revision notes
- <N> starts a source beat; text after the marker is an optional heading
- red font runs are exact emphasis candidates
- <끝> ends the manuscript
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document

MARKER_RE = re.compile(r"^\s*<(\d+)>\s*(.*)$")
END_RE = re.compile(r"^\s*<끝>\s*$")
REVISION_RE = re.compile(r"^\s*P(\d+)\.\s*(.*)$", re.IGNORECASE)
SECTION_RE = re.compile(r"^\s*(\d+)\.\s+(.+)$")


def is_red(run) -> bool:
    color = run.font.color
    return bool(color and color.rgb and str(color.rgb).upper() == "FF0000")


def run_record(run) -> dict:
    return {
        "text": run.text,
        "emphasis": is_red(run),
        "bold": bool(run.bold),
        "italic": bool(run.italic),
    }


def paragraph_record(paragraph, source_paragraph: int) -> dict:
    runs = [run_record(r) for r in paragraph.runs if r.text]
    return {
        "source_paragraph": source_paragraph,
        "text": paragraph.text,
        "runs": runs,
        "emphasis_texts": [r["text"] for r in runs if r["emphasis"] and r["text"].strip()],
    }


def extract(path: Path) -> dict:
    raw = path.read_bytes()
    doc = Document(path)

    revision_notes = []
    preface = []
    blocks = []
    current = None
    saw_first_marker = False
    ended = False
    accounted_nonempty = 0
    revision_continuation_open = False

    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            if not saw_first_marker:
                revision_continuation_open = False
            continue
        accounted_nonempty += 1

        if END_RE.match(text):
            ended = True
            current = None
            continue

        marker = MARKER_RE.match(text)
        if marker:
            saw_first_marker = True
            number = int(marker.group(1))
            heading = marker.group(2).strip()
            section_match = SECTION_RE.match(heading)
            current = {
                "block": number,
                "source_marker_paragraph": idx,
                "heading": heading,
                "is_section_opener": bool(section_match),
                "section_number": int(section_match.group(1)) if section_match else None,
                "section_title": section_match.group(2).strip() if section_match else "",
                "paragraphs": [],
            }
            blocks.append(current)
            continue

        if not saw_first_marker:
            revision = REVISION_RE.match(text)
            if revision:
                revision_notes.append({
                    "display_page": int(revision.group(1)),
                    "instruction": revision.group(2).strip(),
                    "continuation": [],
                    "source_paragraph": idx,
                    "text": paragraph.text,
                })
                revision_continuation_open = True
            elif revision_continuation_open and revision_notes:
                revision_notes[-1]["continuation"].append(paragraph.text)
                revision_notes[-1]["instruction"] += "\n" + paragraph.text.strip()
            else:
                preface.append(paragraph_record(paragraph, idx))
            continue

        if current is None:
            raise ValueError(f"Unaccounted text after manuscript end at paragraph {idx}: {text}")
        current["paragraphs"].append(paragraph_record(paragraph, idx))

    if not saw_first_marker:
        raise ValueError("No <N> source markers found")

    emphasis_count = sum(
        len(p["emphasis_texts"])
        for block in blocks
        for p in block["paragraphs"]
    )

    return {
        "source": {
            "path": str(path.resolve()),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "paragraphs_total": len(doc.paragraphs),
            "nonempty_paragraphs_accounted": accounted_nonempty,
        },
        "preface": preface,
        "revision_notes": revision_notes,
        "blocks": blocks,
        "summary": {
            "block_count": len(blocks),
            "first_block": blocks[0]["block"],
            "last_block": blocks[-1]["block"],
            "section_openers": sum(1 for b in blocks if b["is_section_opener"]),
            "emphasis_run_count": emphasis_count,
            "end_marker_found": ended,
        },
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("-o", "--output", type=Path, required=True)
    args = ap.parse_args()

    data = extract(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(data["summary"], ensure_ascii=False))


if __name__ == "__main__":
    main()
