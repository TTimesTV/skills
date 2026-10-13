"""Convert a reviewed TTimes CMS export to an inline Word archive. No network access."""
from pathlib import Path
import argparse,json,hashlib,re,collections
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX

def text_of(block):
    # Explicit empty finalContent is an approved deletion, not a fallback request.
    for k in ('finalContent','correctedContent','reviewedContent','rawContent'):
        if block.get(k) is not None:return block[k]
    return ''

def offset(text,n):
    # CMS offsets are browser UTF-16 code units, not Python code points.
    raw=text.encode('utf-16-le');assert isinstance(n,int) and 0<=n<=len(raw)//2
    return len(raw[:n*2].decode('utf-16-le'))

def setup(title):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    for el in list(d.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
    s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    n=d.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.line_spacing=1.15;n.paragraph_format.space_after=Pt(4)
    d.styles['Title'].font.size=Pt(21)
    for name,size in [('Heading 1',15),('Heading 2',12),('Heading 3',11)]:d.styles[name].font.size=Pt(size)
    for name in ['Speech','Production','EditorNote']:
        st=d.styles.add_style(name,1);st.base_style=d.styles['Normal']
    d.styles['Production'].paragraph_format.space_after=Pt(2)
    d.styles['EditorNote'].font.size=Pt(9);d.styles['EditorNote'].font.color.rgb=RGBColor.from_string('A12622');d.styles['EditorNote'].paragraph_format.keep_with_next=True
    footer=s.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    d.add_paragraph(title,'Title');d.core_properties.title=title;d.core_properties.author='티타임즈'
    return d

def emit_object(d,x,unlinked=False):
    t=x['text'];kind=x['kind']
    if unlinked:d.add_paragraph('적용 구간 미지정 · 아래 항목은 이 발화 블록에 배치돼 있으나 정확한 시작과 끝은 저장돼 있지 않습니다','EditorNote')
    if t.strip()=='.':
        d.add_paragraph('연결 표식 .  — JSON의 표식을 보존하며 특정 자료의 연속 노출로 단정하지 않음','EditorNote');return
    # The CMS category 'caption' also stores material calls and speech corrections.
    # Do not invent //자막: that would change the spoken-caption suppression contract.
    for line in t.split('\n'):
        p=d.add_paragraph(style='Production');r=p.add_run(line)
        r.bold=True;r.font.color.rgb=RGBColor.from_string('007C91');r.font.highlight_color=WD_COLOR_INDEX.YELLOW
        p.paragraph_format.keep_with_next=True

def kind(t):
    q=t.strip()
    if q=='.' or q.startswith(('(', '번역)', '각 발화')) or '<-' in q:return 'note'
    if re.match(r'^(?://)?\d+(?:_\d+)?\.',q) or q.startswith(('//프로필','//그래픽')):return 'material'
    return 'caption'

def convert(source,output,manual_pending=False):
    data=json.loads(source.read_text(encoding='utf-8-sig'));blocks=sorted(data['script']['blocks'],key=lambda b:b['blockIndex']);bm={b['blockIndex']:b for b in blocks};selected=[];excluded=collections.Counter();issues=[]
    for family in ['highlights','visuals','insertCuts']:
        for c in data[family]['candidates']:
            approved=c.get('status')=='used' or (family=='highlights' and manual_pending and c.get('status')=='pending' and c.get('subtitleType')=='manual')
            if not approved:excluded[(family,c.get('status'))]+=1;continue
            if family!='highlights':raise ValueError('Approved visual/insertCut requires explicit schema adapter, not silent omission')
            bi=next((c[k] for k in ['linkedRangeBlockIndex','placementBlockIndex','blockIndex'] if c.get(k) is not None),None)
            assert bi in bm
            t=c['editedSubtitle'] if c.get('editedSubtitle') is not None else c['subtitle'];speech=text_of(bm[bi]);a=c.get('linkedRangeStart');z=c.get('linkedRangeEnd')
            if a is not None and z is not None:
                a,z=offset(speech,a),offset(speech,z);assert a<z,(c['id'],a,z)
            else:a=z=None;issues.append({'id':c['id'],'block':bi,'issue':'no exact range','text':t})
            selected.append(dict(id=c['id'],bi=bi,text=t,a=a,z=z,order=c.get('order',0),kind=kind(t),status=c['status'],anchor=speech[a:z] if a is not None else None,color=c.get('colorTag')))
    selected.sort(key=lambda x:(x['bi'],x['a'] if x['a'] is not None else -1,x['order'],x['id']))
    d=setup(data['meta']['name']+' 최종 화면 구성')
    stamp=re.search(r'(\d{4})(\d{2})(\d{2})$',source.stem)
    dated=f'{stamp[1]}년 {int(stamp[2])}월 {int(stamp[3])}일 ' if stamp else ''
    d.add_paragraph(f'사용자가 최종 승인본으로 지정한 {dated}JSON을 Word로 복원했습니다. 화면 문구와 발화는 저장된 최종 텍스트를 유지했습니다. 시각은 원고 블록의 시작 시각이며, 문장별 초 단위 타이밍을 새로 추정하지 않았습니다.')
    d.add_paragraph('노란색은 승인된 지시와 문구, 청록색은 정확히 연결된 발화, 검정은 일반 발화입니다. 빨간색은 복원 과정에서 붙인 편집 메모입니다. 적용 구간이 겹치는 곳에는 계속되는 구성과 끝나는 구성을 표시했습니다. 번호와 파일명은 원본 그대로이며 실제 이미지·영상 파일은 JSON에 포함돼 있지 않습니다. 원본에 없는 //자막 표시는 추가하지 않았습니다. 괄호 속 컷·강조·배치 지시는 그대로 보존합니다.')
    if manual_pending:d.add_paragraph(f'이 파일은 직접 등록한 배치가 대부분 pending으로 저장되어 있습니다. 이번 최종본 지정에 따라 used 항목과 폐기되지 않은 직접 등록 pending 항목, 총 {len(selected)}개를 복원했습니다. 자동 제안 전체를 승인한 것으로 해석하지 않았습니다.','EditorNote')
    d.add_paragraph('단독 점 표식은 표시 문구로 출력하지 않고 연결 표식으로 남겼습니다. 구간 미지정 항목은 블록 첫머리에 기록하되 임의로 발화 전체에 적용하지 않았습니다.','EditorNote')
    d.add_heading('승인 원고와 화면 구성',1)
    emitted=[];speech_records=[];overlaps=[]
    for b in blocks:
        bi=b['blockIndex'];s=text_of(b);xs=[x for x in selected if x['bi']==bi]
        d.add_heading(f"{b.get('speaker') or '화자 미지정'} {b.get('startTime') or ''}",2)
        for x in xs:
            if x['a'] is None:emit_object(d,x,True);emitted.append(x['id'])
        if s=='':d.add_paragraph('최종 원고에서 삭제된 블록','EditorNote');continue
        linked=[x for x in xs if x['a'] is not None];points=sorted({0,len(s)}|{x[k] for x in linked for k in ['a','z']});fragments=[]
        for j,a in enumerate(points[:-1]):
            starts=[x for x in linked if x['a']==a];continuing=[x for x in linked if x['a']<a<x['z']];ended=[x for x in linked if x['z']==a]
            if ended and continuing:
                d.add_paragraph('종료: '+' / '.join(x['text'].split('\n')[0][:45] for x in ended)+' · 나머지 구성 계속','EditorNote')
                overlaps.append({'block':bi,'position':a,'ended':[x['id'] for x in ended],'continuing':[x['id'] for x in continuing]})
            if starts and continuing:d.add_paragraph('앞 구성 계속 · 다음 항목을 함께 적용','EditorNote')
            for x in starts:emit_object(d,x);emitted.append(x['id'])
            z=points[j+1];fragment=s[a:z];p=d.add_paragraph(style='Speech');r=p.add_run(fragment)
            active=[x for x in linked if x['a']<=a<x['z']]
            if active:r.font.color.rgb=RGBColor.from_string('007C91')
            p.paragraph_format.keep_with_next=False
            fragments.append(fragment)
            speech_records.append({'block':bi,'a':a,'z':z,'text':fragment,'active':[x['id'] for x in active]})
        assert ''.join(fragments)==s
    revisions=(data.get('revision',{}).get('revision') or {}).get('items',[])
    if revisions:
        d.add_page_break();d.add_heading('후속 영상 수정 기록',1)
        d.add_paragraph('이 시각은 편집 영상 기준입니다. 앞의 원고 시각과 다른 시간축이므로 원고에 자동 합치지 않았습니다. 완료 표시는 JSON에 기록된 처리 상태이며, 영상에 실제 반영됐는지는 이번 변환에서 재검증하지 않았습니다.')
        for rev in sorted(revisions,key=lambda x:x['startTimeSec']):
            t=rev['startTimeSec'];m,sec=divmod(t,60);label=f'{int(m):02d}:{int(sec):02d}';status={'complete':'완료','pending':'미완료'}.get(rev['status'],rev['status'])
            d.add_heading(label+' '+status,3);d.add_paragraph(rev['content'].replace('\r\n','\n').replace('\r','\n'))
            if rev.get('screenshotUrl'):d.add_paragraph('수정 스크린샷은 서버 참조만 저장돼 있어 이 파일에는 포함하지 않았습니다.','EditorNote')
    # Verify the saved visible transcript, not only the in-memory source.
    output.parent.mkdir(exist_ok=True,parents=True);d.save(output);rd=Document(output)
    visible=[p.text for p in rd.paragraphs if p.style.name=='Speech'];assert visible==[x['text'] for x in speech_records]
    assert sorted(emitted)==sorted(x['id'] for x in selected)
    report={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'output':str(output),'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'blocks':len(blocks),'empty_final_blocks':[b['blockIndex'] for b in blocks if text_of(b)==''],'selected_count':len(selected),'selected_status':dict(collections.Counter(x['status'] for x in selected)),'kind_counts':dict(collections.Counter(x['kind'] for x in selected)),'excluded':{str(k):v for k,v in excluded.items()},'issues':issues,'overlap_events':overlaps,'revision_count':len(revisions),'not_in_screen_document':['unselected automatic suggestions','highlightReel and contentSet marketing variants; not screen composition'],'speech_roundtrip':'pass','selected_coverage':'pass','visual_review':'not_run','selected':selected,'speech_fragments':speech_records}
    return report

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('source',type=Path);ap.add_argument('output',type=Path);ap.add_argument('--include-manual-pending',action='store_true');ap.add_argument('--audit',type=Path,required=True);args=ap.parse_args()
    r=convert(args.source,args.output,args.include_manual_pending);args.audit.write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8');print({k:r[k] for k in ['blocks','selected_count','selected_status','kind_counts','revision_count','speech_roundtrip','selected_coverage']})
